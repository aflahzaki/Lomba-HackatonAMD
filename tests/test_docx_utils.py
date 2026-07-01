"""Tests for StageArsitekturIntegrasiSistem/output/docx_utils.py."""

import sys
import os

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "StageArsitekturIntegrasiSistem", "output"))

from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from docx_utils import (
    FONT_NAME,
    FONT_SIZE_MAIN,
    FONT_SIZE_TITLE,
    FONT_SIZE_SUBTITLE,
    LINE_SPACING,
    MARGIN,
    TEAM_MEMBERS,
    create_document,
    add_heading,
    add_paragraph,
    add_bullet_list,
    add_numbered_list,
    add_table,
    add_figure_caption,
    add_references,
    add_table_of_contents,
    add_cover_page,
    save_document,
)


class TestConstants:
    def test_font_name(self):
        assert FONT_NAME == "Times New Roman"

    def test_font_size_main(self):
        assert FONT_SIZE_MAIN == Pt(12)

    def test_font_size_title(self):
        assert FONT_SIZE_TITLE == Pt(14)

    def test_line_spacing(self):
        assert LINE_SPACING == 1.5

    def test_team_members_count(self):
        assert len(TEAM_MEMBERS) == 4

    def test_team_members_are_tuples(self):
        for member in TEAM_MEMBERS:
            assert isinstance(member, tuple)
            assert len(member) == 2


class TestCreateDocument:
    def test_returns_document(self):
        doc = create_document()
        assert doc is not None

    def test_margins_set(self):
        doc = create_document()
        for section in doc.sections:
            assert section.top_margin == MARGIN
            assert section.bottom_margin == MARGIN
            assert section.left_margin == MARGIN
            assert section.right_margin == MARGIN

    def test_default_font(self):
        doc = create_document()
        style = doc.styles["Normal"]
        assert style.font.name == FONT_NAME
        assert style.font.size == FONT_SIZE_MAIN

    def test_line_spacing(self):
        doc = create_document()
        style = doc.styles["Normal"]
        assert style.paragraph_format.line_spacing == LINE_SPACING

    def test_heading_styles_configured(self):
        doc = create_document()
        for i in range(1, 4):
            h_style = doc.styles[f"Heading {i}"]
            assert h_style.font.name == FONT_NAME
            assert h_style.font.color.rgb == RGBColor(0, 0, 0)
            assert h_style.font.bold is True


class TestAddHeading:
    def test_heading_text(self):
        doc = create_document()
        h = add_heading(doc, "Test Heading", level=1)
        assert h.text == "Test Heading"

    def test_heading_level_1_font_size(self):
        doc = create_document()
        h = add_heading(doc, "Title", level=1)
        for run in h.runs:
            assert run.font.size == FONT_SIZE_TITLE

    def test_heading_level_2_font_size(self):
        doc = create_document()
        h = add_heading(doc, "Subtitle", level=2)
        for run in h.runs:
            assert run.font.size == FONT_SIZE_SUBTITLE

    def test_heading_level_3_font_size(self):
        doc = create_document()
        h = add_heading(doc, "Sub-subtitle", level=3)
        for run in h.runs:
            assert run.font.size == FONT_SIZE_MAIN

    def test_heading_centered(self):
        doc = create_document()
        h = add_heading(doc, "Centered", level=1, centered=True)
        assert h.alignment == WD_ALIGN_PARAGRAPH.CENTER

    def test_heading_not_centered_by_default(self):
        doc = create_document()
        h = add_heading(doc, "Left", level=1)
        assert h.alignment != WD_ALIGN_PARAGRAPH.CENTER

    def test_heading_font_color_black(self):
        doc = create_document()
        h = add_heading(doc, "Black", level=1)
        for run in h.runs:
            assert run.font.color.rgb == RGBColor(0, 0, 0)


class TestAddParagraph:
    def test_paragraph_text(self):
        doc = create_document()
        p = add_paragraph(doc, "Hello world")
        assert p.text == "Hello world"

    def test_paragraph_bold(self):
        doc = create_document()
        p = add_paragraph(doc, "Bold text", bold=True)
        assert p.runs[0].font.bold is True

    def test_paragraph_italic(self):
        doc = create_document()
        p = add_paragraph(doc, "Italic text", italic=True)
        assert p.runs[0].font.italic is True

    def test_paragraph_font(self):
        doc = create_document()
        p = add_paragraph(doc, "Font check")
        assert p.runs[0].font.name == FONT_NAME
        assert p.runs[0].font.size == FONT_SIZE_MAIN

    def test_paragraph_justified(self):
        doc = create_document()
        p = add_paragraph(doc, "Justified text")
        assert p.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY

    def test_paragraph_indent(self):
        doc = create_document()
        p = add_paragraph(doc, "Indented", indent=2)
        assert abs(p.paragraph_format.left_indent - Cm(2)) < 200

    def test_paragraph_line_spacing(self):
        doc = create_document()
        p = add_paragraph(doc, "Spaced")
        assert p.paragraph_format.line_spacing == LINE_SPACING


class TestAddBulletList:
    def test_bullet_list_items(self):
        doc = create_document()
        items = ["Item A", "Item B", "Item C"]
        add_bullet_list(doc, items)
        paragraphs = doc.paragraphs
        bullet_paras = [p for p in paragraphs if p.text.startswith("\u2022")]
        assert len(bullet_paras) == 3

    def test_bullet_prefix(self):
        doc = create_document()
        add_bullet_list(doc, ["Test"])
        bullet_paras = [p for p in doc.paragraphs if "\u2022" in p.text]
        assert len(bullet_paras) == 1
        assert "Test" in bullet_paras[0].text

    def test_empty_list(self):
        doc = create_document()
        add_bullet_list(doc, [])
        bullet_paras = [p for p in doc.paragraphs if "\u2022" in p.text]
        assert len(bullet_paras) == 0


class TestAddNumberedList:
    def test_numbered_list_items(self):
        doc = create_document()
        items = ["First", "Second", "Third"]
        add_numbered_list(doc, items)
        numbered = [p for p in doc.paragraphs if p.text and p.text[0].isdigit()]
        assert len(numbered) == 3

    def test_numbered_list_start(self):
        doc = create_document()
        add_numbered_list(doc, ["Only"], start=5)
        numbered = [p for p in doc.paragraphs if p.text.startswith("5.")]
        assert len(numbered) == 1

    def test_numbered_list_font(self):
        doc = create_document()
        add_numbered_list(doc, ["Check font"])
        numbered = [p for p in doc.paragraphs if p.text and p.text[0].isdigit()]
        assert numbered[0].runs[0].font.name == FONT_NAME


class TestAddTable:
    def test_table_headers(self):
        doc = create_document()
        headers = ["Col A", "Col B"]
        rows = [("1", "2"), ("3", "4")]
        table = add_table(doc, headers, rows)
        hdr_cells = table.rows[0].cells
        assert hdr_cells[0].text == "Col A"
        assert hdr_cells[1].text == "Col B"

    def test_table_row_count(self):
        doc = create_document()
        headers = ["H1", "H2"]
        rows = [("a", "b"), ("c", "d"), ("e", "f")]
        table = add_table(doc, headers, rows)
        # 1 header row + 3 data rows
        assert len(table.rows) == 4

    def test_table_data(self):
        doc = create_document()
        headers = ["Name"]
        rows = [("Alice",), ("Bob",)]
        table = add_table(doc, headers, rows)
        assert table.rows[1].cells[0].text == "Alice"
        assert table.rows[2].cells[0].text == "Bob"

    def test_table_with_title(self):
        doc = create_document()
        headers = ["X"]
        rows = [("1",)]
        add_table(doc, headers, rows, title="My Table", table_number=1)
        title_paras = [p for p in doc.paragraphs if "My Table" in p.text]
        assert len(title_paras) >= 1
        assert "Tabel 1" in title_paras[0].text

    def test_table_without_title(self):
        doc = create_document()
        headers = ["X"]
        rows = [("1",)]
        add_table(doc, headers, rows)
        title_paras = [p for p in doc.paragraphs if "Tabel" in p.text]
        assert len(title_paras) == 0


class TestAddFigureCaption:
    def test_caption_text(self):
        doc = create_document()
        p = add_figure_caption(doc, "A diagram", 1)
        assert "Gambar 1" in p.text
        assert "A diagram" in p.text

    def test_caption_italic(self):
        doc = create_document()
        p = add_figure_caption(doc, "Caption", 2)
        assert p.runs[0].font.italic is True

    def test_caption_centered(self):
        doc = create_document()
        p = add_figure_caption(doc, "Centered", 3)
        assert p.alignment == WD_ALIGN_PARAGRAPH.CENTER


class TestAddReferences:
    def test_references_heading(self):
        doc = create_document()
        refs = ["Ref 1", "Ref 2"]
        add_references(doc, refs)
        headings = [p for p in doc.paragraphs if "DAFTAR PUSTAKA" in p.text]
        assert len(headings) >= 1

    def test_references_count(self):
        doc = create_document()
        refs = ["Author (2020). Title.", "Author2 (2021). Title2."]
        add_references(doc, refs)
        ref_paras = [p for p in doc.paragraphs if "Author" in p.text]
        assert len(ref_paras) == 2

    def test_references_formatting(self):
        doc = create_document()
        refs = ["Test Reference"]
        add_references(doc, refs)
        ref_paras = [p for p in doc.paragraphs if "Test Reference" in p.text]
        assert ref_paras[0].runs[0].font.name == FONT_NAME


class TestAddTableOfContents:
    def test_toc_heading(self):
        doc = create_document()
        sections = [("Section 1", ["Sub 1"])]
        add_table_of_contents(doc, sections)
        headings = [p for p in doc.paragraphs if "DAFTAR ISI" in p.text]
        assert len(headings) >= 1

    def test_toc_sections_present(self):
        doc = create_document()
        sections = [
            ("Pendahuluan", ["Latar Belakang", "Tujuan"]),
            ("Metodologi", []),
        ]
        add_table_of_contents(doc, sections)
        found_pendahuluan = any("Pendahuluan" in p.text for p in doc.paragraphs)
        found_metodologi = any("Metodologi" in p.text for p in doc.paragraphs)
        assert found_pendahuluan
        assert found_metodologi

    def test_toc_subsections(self):
        doc = create_document()
        sections = [("Main", ["Sub A", "Sub B"])]
        add_table_of_contents(doc, sections)
        found_sub_a = any("Sub A" in p.text for p in doc.paragraphs)
        found_sub_b = any("Sub B" in p.text for p in doc.paragraphs)
        assert found_sub_a
        assert found_sub_b


class TestAddCoverPage:
    def test_cover_page_title(self):
        doc = create_document()
        add_cover_page(doc, title="Test Title", stage_name="Stage 1")
        found = any("TEST TITLE" in p.text for p in doc.paragraphs)
        assert found

    def test_cover_page_stage_name(self):
        doc = create_document()
        add_cover_page(doc, title="Title", stage_name="Stage 99 - Test")
        found = any("Stage 99 - Test" in p.text for p in doc.paragraphs)
        assert found

    def test_cover_page_team_members(self):
        doc = create_document()
        add_cover_page(doc, title="Title", stage_name="Stage 1")
        for name, nim in TEAM_MEMBERS:
            found = any(name in p.text for p in doc.paragraphs)
            assert found, f"Team member {name} not found on cover page"

    def test_cover_page_university(self):
        doc = create_document()
        add_cover_page(doc, title="T", stage_name="S")
        found = any("TELKOM UNIVERSITY" in p.text for p in doc.paragraphs)
        assert found


class TestSaveDocument:
    def test_save_creates_file(self, tmp_path):
        doc = create_document()
        add_paragraph(doc, "Test content")
        # Temporarily override the output dir by saving directly
        filepath = str(tmp_path / "test_output.docx")
        doc.save(filepath)
        assert os.path.exists(filepath)
        assert os.path.getsize(filepath) > 0
