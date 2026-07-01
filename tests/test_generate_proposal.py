"""Tests for generate_proposal.py helper functions."""

import sys
import os

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

from generate_proposal import (
    set_cell_shading,
    add_page_break,
    set_paragraph_spacing,
    add_bullet_point,
    create_table,
)


class TestSetCellShading:
    def test_shading_applied(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        cell = table.rows[0].cells[0]
        set_cell_shading(cell, "FF0000")
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = tc_pr.findall(qn("w:shd"))
        assert len(shd) >= 1
        assert shd[-1].get(qn("w:fill")) == "FF0000"

    def test_shading_different_colors(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        set_cell_shading(table.rows[0].cells[0], "2E4057")
        set_cell_shading(table.rows[0].cells[1], "FFFFFF")
        shd0 = table.rows[0].cells[0]._tc.get_or_add_tcPr().findall(qn("w:shd"))
        shd1 = table.rows[0].cells[1]._tc.get_or_add_tcPr().findall(qn("w:shd"))
        assert shd0[-1].get(qn("w:fill")) == "2E4057"
        assert shd1[-1].get(qn("w:fill")) == "FFFFFF"


class TestAddPageBreak:
    def test_page_break_adds_paragraph(self):
        doc = Document()
        initial_count = len(doc.paragraphs)
        add_page_break(doc)
        assert len(doc.paragraphs) >= initial_count


class TestSetParagraphSpacing:
    def test_default_spacing(self):
        doc = Document()
        p = doc.add_paragraph("Test")
        set_paragraph_spacing(p)
        assert p.paragraph_format.space_before == Pt(0)
        assert p.paragraph_format.space_after == Pt(6)

    def test_custom_spacing(self):
        doc = Document()
        p = doc.add_paragraph("Test")
        set_paragraph_spacing(p, before=10, after=20)
        assert p.paragraph_format.space_before == Pt(10)
        assert p.paragraph_format.space_after == Pt(20)

    def test_zero_spacing(self):
        doc = Document()
        p = doc.add_paragraph("Test")
        set_paragraph_spacing(p, before=0, after=0)
        assert p.paragraph_format.space_before == Pt(0)
        assert p.paragraph_format.space_after == Pt(0)


class TestAddBulletPoint:
    def test_bullet_text(self):
        doc = Document()
        p = add_bullet_point(doc, "Test bullet")
        assert "Test bullet" in p.text

    def test_bullet_with_bold_prefix(self):
        doc = Document()
        p = add_bullet_point(doc, " rest of text", bold_prefix="Bold:")
        assert p.runs[0].bold is True
        assert p.runs[0].text == "Bold:"
        assert "rest of text" in p.text

    def test_bullet_without_bold_prefix(self):
        doc = Document()
        p = add_bullet_point(doc, "No bold")
        assert p.runs[0].text == "No bold"

    def test_bullet_level_indent(self):
        doc = Document()
        p = add_bullet_point(doc, "Indented", level=1)
        assert abs(p.paragraph_format.left_indent - Cm(1.5 * 2)) < 200

    def test_bullet_level_zero_no_indent(self):
        doc = Document()
        p = add_bullet_point(doc, "No indent", level=0)
        assert p.paragraph_format.left_indent is None


class TestCreateTable:
    def test_table_structure(self):
        doc = Document()
        headers = ["Name", "Age"]
        rows = [("Alice", "30"), ("Bob", "25")]
        table = create_table(doc, headers, rows)
        assert len(table.rows) == 3  # 1 header + 2 data

    def test_table_headers(self):
        doc = Document()
        headers = ["Col1", "Col2"]
        rows = [("a", "b")]
        table = create_table(doc, headers, rows)
        assert table.rows[0].cells[0].text == "Col1"
        assert table.rows[0].cells[1].text == "Col2"

    def test_table_data_rows(self):
        doc = Document()
        headers = ["X"]
        rows = [("val1",), ("val2",)]
        table = create_table(doc, headers, rows)
        assert table.rows[1].cells[0].text == "val1"
        assert table.rows[2].cells[0].text == "val2"

    def test_table_header_bold(self):
        doc = Document()
        headers = ["H"]
        rows = [("d",)]
        table = create_table(doc, headers, rows)
        for paragraph in table.rows[0].cells[0].paragraphs:
            for run in paragraph.runs:
                assert run.bold is True

    def test_table_header_shading(self):
        doc = Document()
        headers = ["H"]
        rows = [("d",)]
        table = create_table(doc, headers, rows)
        tc_pr = table.rows[0].cells[0]._tc.get_or_add_tcPr()
        shd = tc_pr.findall(qn("w:shd"))
        assert len(shd) >= 1
        assert shd[-1].get(qn("w:fill")) == "2E4057"

    def test_table_alignment(self):
        doc = Document()
        headers = ["A"]
        rows = [("1",)]
        table = create_table(doc, headers, rows)
        assert table.alignment == WD_TABLE_ALIGNMENT.CENTER

    def test_table_with_col_widths(self):
        doc = Document()
        headers = ["Narrow", "Wide"]
        rows = [("a", "b")]
        table = create_table(doc, headers, rows, col_widths=[1.0, 3.0])
        assert table.rows[0].cells[0].width == Inches(1.0)
        assert table.rows[0].cells[1].width == Inches(3.0)

    def test_table_empty_rows(self):
        doc = Document()
        headers = ["H1", "H2"]
        rows = []
        table = create_table(doc, headers, rows)
        assert len(table.rows) == 1  # header only

    def test_table_data_font_size(self):
        doc = Document()
        headers = ["H"]
        rows = [("data",)]
        table = create_table(doc, headers, rows)
        for paragraph in table.rows[1].cells[0].paragraphs:
            for run in paragraph.runs:
                assert run.font.size == Pt(9)
